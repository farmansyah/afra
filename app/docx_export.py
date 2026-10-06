"""Markdown -> formatted DOCX with academic templates and automatic references."""
import io
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .citations import STYLES, bibliography, process_citations

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "right": WD_ALIGN_PARAGRAPH.RIGHT}

LABELS = {
    "en": {"abstract": "Abstract", "keywords": "Keywords", "references": "References", "toc": "Table of Contents",
           "chapter": "CHAPTER", "index_terms": "Index Terms"},
    "id": {"abstract": "Abstrak", "keywords": "Kata kunci", "references": "Daftar Pustaka", "toc": "Daftar Isi",
           "chapter": "BAB", "index_terms": "Kata Kunci"},
}

TEMPLATES = {
    "apa7": {
        "name": "APA 7th edition (student/professional paper)", "page": "letter", "margins": (2.54, 2.54, 2.54, 2.54),
        "font": "Times New Roman", "size": 12, "spacing": 2.0, "after": 0, "indent": 1.27, "align": "left",
        "numbering": None, "title_page": "apa", "columns": 1, "style": "apa", "page_number": "header-right",
        "h": {1: dict(align="center", bold=True), 2: dict(align="left", bold=True),
              3: dict(align="left", bold=True, italic=True), 4: dict(align="left", bold=True)},
        "ref_heading": dict(align="center", bold=True), "ref_size": None, "ref_spacing": None,
    },
    "ieee": {
        "name": "IEEE conference/journal (two-column)", "page": "letter", "margins": (1.9, 2.54, 1.57, 1.57),
        "font": "Times New Roman", "size": 10, "spacing": 1.0, "after": 0, "indent": 0.36, "align": "justify",
        "numbering": "ieee", "title_page": "ieee", "columns": 2, "style": "ieee", "page_number": None,
        "h": {1: dict(align="center", small_caps=True, size=10, before=8, after=4),
              2: dict(align="left", italic=True, size=10, before=6, after=3),
              3: dict(align="left", italic=True, size=10)},
        "ref_heading": dict(align="center", small_caps=True, size=10), "ref_size": 8, "ref_spacing": 1.0,
    },
    "journal": {
        "name": "Generic journal article (Elsevier/Springer-like)", "page": "a4", "margins": (2.5, 2.5, 2.5, 2.5),
        "font": "Times New Roman", "size": 11, "spacing": 1.15, "after": 6, "indent": 0, "align": "justify",
        "numbering": "decimal", "title_page": "journal", "columns": 1, "style": "apa", "page_number": "footer-center",
        "h": {1: dict(align="left", bold=True, size=12, before=12, after=6),
              2: dict(align="left", bold=True, italic=True, before=10, after=4),
              3: dict(align="left", italic=True, before=8, after=4)},
        "ref_heading": dict(align="left", bold=True, size=12), "ref_size": 10, "ref_spacing": 1.0,
    },
    "thesis": {
        "name": "Thesis / Skripsi (Indonesian university format)", "page": "a4", "margins": (4, 3, 4, 3),
        "font": "Times New Roman", "size": 12, "spacing": 1.5, "after": 0, "indent": 1.25, "align": "justify",
        "numbering": "chapter", "title_page": "thesis", "columns": 1, "style": "apa", "page_number": "footer-center",
        "h": {1: dict(align="center", bold=True, upper=True, page_break=True, after=18),
              2: dict(align="left", bold=True, before=12, after=6),
              3: dict(align="left", bold=True, before=6, after=6)},
        "ref_heading": dict(align="center", bold=True, upper=True, page_break=True), "ref_size": None, "ref_spacing": 1.0,
    },
    "report": {
        "name": "Research report (clean, numbered headings)", "page": "a4", "margins": (2.5, 2.5, 2.5, 2.5),
        "font": "Calibri", "size": 11, "spacing": 1.15, "after": 8, "indent": 0, "align": "justify",
        "numbering": "decimal", "title_page": "report", "columns": 1, "style": "apa", "page_number": "footer-center",
        "h": {1: dict(align="left", bold=True, size=16, before=18, after=8, color="1F3864"),
              2: dict(align="left", bold=True, size=13, before=12, after=6, color="1F3864"),
              3: dict(align="left", bold=True, italic=True, size=11, before=8, after=4)},
        "ref_heading": dict(align="left", bold=True, size=16, color="1F3864"), "ref_size": 10, "ref_spacing": 1.0,
    },
}

PAGE_SIZES = {"a4": (21.0, 29.7), "letter": (21.59, 27.94)}
REF_HEADINGS = re.compile(r"^(references|reference list|bibliography|works cited|daftar pustaka|referensi|literature cited)$", re.I)
ABSTRACT_HEADINGS = re.compile(r"^(abstract|abstrak|summary)$", re.I)
NUM_PREFIX = re.compile(r"^((bab|chapter)\s+[ivxlc\d]+[.:]?\s*|[ivxlc]+\.\s+|[A-Z]\.\s+|\d+(\.\d+)*\.?\s+)", re.I)


# ---------------------------------------------------------------- low-level helpers
def _set_run_font(run, name=None, size=None, bold=None, italic=None, small_caps=None, color=None):
    f = run.font
    if name:
        f.name = name
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), name)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if small_caps is not None:
        f.small_caps = small_caps
    if color:
        f.color.rgb = RGBColor.from_string(color)


def _style_font(style, name, size, bold=False, italic=False, color="000000"):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        if attr.endswith("Theme"):
            del rfonts.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _add_field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    ph = OxmlElement("w:t"); ph.text = placeholder
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, t, s, ph, e):
        run._r.append(el)
    return run


def _set_columns(section, num):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), "360")


def _roman(n):
    vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
            (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


# ---------------------------------------------------------------- inline markdown
_INLINE = re.compile(
    r"(\*\*\*[^*]+?\*\*\*|\*\*[^*]+?\*\*|__[^_]+?__|(?<![\w*])\*(?!\s)[^*]+?(?<!\s)\*(?!\w)|(?<!\w)_(?!\s)[^_]+?(?<!\s)_(?!\w)"
    r"|`[^`]+`|\[[^\]]+\]\([^)]+\)|\$[^$\n]+\$|\^[^^\s]+\^|~[^~\s]+~|<sup>.*?</sup>|<sub>.*?</sub>)")


def add_inline(paragraph, text, base: dict | None = None):
    base = base or {}
    text = text.replace("\\*", "\u0000").replace("\\_", "\u0001")
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos:m.start()], base)
        tok = m.group(0)
        if tok.startswith("***"):
            _run(paragraph, tok[3:-3], {**base, "bold": True, "italic": True})
        elif tok.startswith("**") or tok.startswith("__"):
            add_inline(paragraph, tok[2:-2], {**base, "bold": True})
        elif tok[0] in "*_":
            add_inline(paragraph, tok[1:-1], {**base, "italic": True})
        elif tok[0] == "`":
            _run(paragraph, tok[1:-1], {**base, "font": "Consolas"})
        elif tok[0] == "[":
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            _run(paragraph, lm.group(1), base)
        elif tok[0] == "$":
            _run(paragraph, tok[1:-1], {**base, "italic": True})
        elif tok[0] == "^" or tok.startswith("<sup>"):
            _run(paragraph, re.sub(r"</?sup>|\^", "", tok), {**base, "sup": True})
        elif tok[0] == "~" or tok.startswith("<sub>"):
            _run(paragraph, re.sub(r"</?sub>|~", "", tok), {**base, "sub": True})
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], base)


def _run(paragraph, text, fmt):
    text = text.replace("\u0000", "*").replace("\u0001", "_")
    if not text:
        return
    r = paragraph.add_run(text)
    if fmt.get("bold"):
        r.bold = True
    if fmt.get("italic"):
        r.italic = True
    if fmt.get("font"):
        _set_run_font(r, name=fmt["font"], size=fmt.get("size"))
    elif fmt.get("size"):
        r.font.size = Pt(fmt["size"])
    if fmt.get("sup"):
        r.font.superscript = True
    if fmt.get("sub"):
        r.font.subscript = True
    if fmt.get("small_caps"):
        r.font.small_caps = True
    if fmt.get("color"):
        r.font.color.rgb = RGBColor.from_string(fmt["color"])
    return r


# ---------------------------------------------------------------- markdown blocks
def parse_blocks(md: str):
    lines = md.replace("\r\n", "\n").split("\n")
    i, blocks = 0, []
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            j = i + 1
            code = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                code.append(lines[j]); j += 1
            blocks.append(("code", "\n".join(code)))
            i = j + 1
            continue
        if s.startswith("$$"):
            if s.endswith("$$") and len(s) > 4:
                blocks.append(("math", s[2:-2])); i += 1; continue
            j = i + 1; eq = []
            while j < len(lines) and not lines[j].strip().endswith("$$"):
                eq.append(lines[j]); j += 1
            if j < len(lines):
                eq.append(lines[j].strip()[:-2])
            blocks.append(("math", " ".join(eq).strip())); i = j + 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*?)\s*#*$", s)
        if m:
            blocks.append(("heading", len(m.group(1)), m.group(2))); i += 1
            continue
        if re.match(r"^(\\pagebreak|\\newpage|<!--\s*pagebreak\s*-->)$", s, re.I):
            blocks.append(("pagebreak",)); i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{2,}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", rows))
            continue
        if re.match(r"^!\[[^\]]*\]\([^)]+\)$", s):
            alt = re.match(r"^!\[([^\]]*)\]", s).group(1)
            blocks.append(("figure", alt)); i += 1
            continue
        if s.startswith(">"):
            q = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q.append(lines[i].strip()[1:].strip()); i += 1
            blocks.append(("quote", " ".join(q)))
            continue
        lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if lm:
            items = []
            while i < len(lines):
                lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", lines[i])
                if lm:
                    items.append([len(lm.group(1).replace("\t", "    ")) // 2, not lm.group(2)[0] in "-*+", lm.group(3)])
                    i += 1
                elif lines[i].strip() and items and lines[i].startswith((" ", "\t")):
                    items[-1][2] += " " + lines[i].strip(); i += 1
                else:
                    break
            blocks.append(("list", items))
            continue
        para = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^\s*(#{1,6}\s|```|\$\$|>|\||[-*+]\s|\d+[.)]\s|!\[)", lines[i]):
            para.append(lines[i].strip()); i += 1
        blocks.append(("para", " ".join(para)))
    return blocks


def _strip_sections(blocks, pattern):
    """Remove headed sections whose heading matches pattern. Returns (blocks, removed_text)."""
    out, removed, skip_level = [], [], None
    for b in blocks:
        if b[0] == "heading":
            if skip_level is not None and b[1] <= skip_level:
                skip_level = None
            if skip_level is None and pattern.match(b[2].strip().strip("*").strip()):
                skip_level = b[1]
                continue
        if skip_level is not None:
            if b[0] == "para":
                removed.append(b[1])
            continue
        out.append(b)
    return out, "\n\n".join(removed)


# ---------------------------------------------------------------- builder
class Builder:
    def __init__(self, tpl: dict, opts: dict):
        self.t = tpl
        self.o = opts
        self.doc = Document()
        self.lang = opts.get("lang", "en")
        self.L = LABELS.get(self.lang, LABELS["en"])
        self.counters = [0] * 7
        self.font = opts.get("font") or tpl["font"]
        self.size = float(opts.get("size") or tpl["size"])
        self.spacing = float(opts.get("spacing") or tpl["spacing"])
        self._setup()

    def _setup(self):
        t, d = self.t, self.doc
        sec = d.sections[0]
        w, h = PAGE_SIZES[self.o.get("page") or t["page"]]
        sec.page_width, sec.page_height = Cm(w), Cm(h)
        top, bottom, left, right = t["margins"]
        sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(top), Cm(bottom), Cm(left), Cm(right)
        normal = d.styles["Normal"]
        _style_font(normal, self.font, self.size)
        pf = normal.paragraph_format
        pf.line_spacing = self.spacing
        pf.space_after = Pt(t["after"])
        pf.space_before = Pt(0)
        for lvl in range(1, 5):
            st = d.styles[f"Heading {lvl}"]
            spec = t["h"].get(lvl) or t["h"].get(3) or {}
            _style_font(st, self.font, spec.get("size") or self.size, bool(spec.get("bold")),
                        bool(spec.get("italic")), spec.get("color") or "000000")
            st.font.small_caps = bool(spec.get("small_caps"))
            hp = st.paragraph_format
            hp.alignment = ALIGN[spec.get("align", "left")]
            hp.space_before = Pt(spec.get("before", 0 if self.spacing >= 2 else 6))
            hp.space_after = Pt(spec.get("after", 0 if self.spacing >= 2 else 3))
            hp.line_spacing = self.spacing
            hp.first_line_indent = Cm(0)
            hp.keep_with_next = True
            hp.page_break_before = False
        for name in ("List Bullet", "List Bullet 2", "Title", "Caption"):
            try:
                _style_font(d.styles[name], self.font, self.size)
            except KeyError:
                pass
        settings = d.settings.element
        upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
        settings.append(upd)
        self._page_numbers(sec)

    def _page_numbers(self, sec):
        pos = self.t.get("page_number")
        if not pos:
            return
        if self.t["title_page"] == "thesis":
            sec.different_first_page_header_footer = True  # no number on the cover
        part = sec.header if pos.startswith("header") else sec.footer
        p = part.paragraphs[0]
        p.alignment = ALIGN["right" if pos.endswith("right") else "center"]
        r = _add_field(p, "PAGE", "1")
        _set_run_font(r, self.font, self.size)

    # -------------------------------------------------------- paragraphs
    def para(self, text, align=None, indent=True, size=None, spacing=None, after=None, base=None, style=None):
        p = self.doc.add_paragraph(style=style)
        pf = p.paragraph_format
        pf.alignment = ALIGN[align or self.t["align"]]
        pf.first_line_indent = Cm(self.t["indent"]) if indent and self.t["indent"] else Cm(0)
        if spacing:
            pf.line_spacing = spacing
        if after is not None:
            pf.space_after = Pt(after)
        b = dict(base or {})
        if size:
            b["size"] = size
        add_inline(p, text, b)
        return p

    def blank(self, n=1):
        for _ in range(n):
            self.doc.add_paragraph()

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def heading(self, level, text):
        level = min(level, 4)
        t = self.t
        spec = t["h"].get(level) or t["h"].get(3) or {}
        numbering = self.o.get("numbering", t["numbering"])
        if numbering:
            text = NUM_PREFIX.sub("", text).strip()
        self.counters[level] += 1
        for k in range(level + 1, 7):
            self.counters[k] = 0
        prefix = ""
        if numbering == "ieee":
            if level == 1:
                prefix = f"{_roman(self.counters[1])}. "
            elif level == 2:
                prefix = f"{chr(64 + self.counters[2])}. "
            elif level == 3:
                prefix = f"{self.counters[3]}) "
        elif numbering in ("decimal", "chapter") and level <= 3:
            nums = ".".join(str(self.counters[k]) for k in range(1, level + 1))
            prefix = f"{nums}. " if (numbering == "decimal" and level == 1) else f"{nums} "
        p = self.doc.add_paragraph(style=f"Heading {level}")
        if spec.get("page_break"):
            p.paragraph_format.page_break_before = True
        if numbering == "chapter" and level == 1:
            r = p.add_run(f"{self.L['chapter']} {_roman(self.counters[1])}")
            r.add_break()
            add_inline(p, text.upper() if spec.get("upper") else text)
            return p
        if numbering == "ieee" and level == 1:
            text = text.upper() if not spec.get("small_caps") else text
        if spec.get("upper"):
            text = text.upper()
        if prefix:
            p.add_run(prefix)
        add_inline(p, text)
        return p

    def list_items(self, items):
        counters = {}
        for depth, ordered, text in items:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = ALIGN[self.t["align"]]
            pf.left_indent = Cm(0.9 + 0.75 * depth)
            pf.first_line_indent = Cm(-0.6)
            pf.space_after = Pt(0)
            if ordered:
                counters[depth] = counters.get(depth, 0) + 1
                for d in list(counters):
                    if d > depth:
                        del counters[d]
                p.add_run(f"{counters[depth]}.\t")
            else:
                p.add_run("•\t" if depth == 0 else "◦\t")
            pf.tab_stops.add_tab_stop(Cm(0.9 + 0.75 * depth))
            add_inline(p, text)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(self.t["after"] or 6)

    def table(self, rows):
        if not rows:
            return
        ncols = max(len(r) for r in rows)
        tbl = self.doc.add_table(rows=len(rows), cols=ncols)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tsize = max(8, self.size - 1)
        for ri, row in enumerate(rows):
            for ci in range(ncols):
                cell = tbl.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                add_inline(p, row[ci] if ci < len(row) else "", {"size": tsize, "bold": ri == 0})
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def body(self, blocks):
        for b in blocks:
            kind = b[0]
            if kind == "heading":
                self.heading(b[1], b[2])
            elif kind == "para":
                cap = re.match(r"^\*{0,2}(Table|Figure|Tabel|Gambar)\s+\d+", b[1])
                if cap:
                    self.para(b[1], align="center" if self.t["columns"] == 1 else "left", indent=False,
                              size=max(8, self.size - 1))
                else:
                    self.para(b[1])
            elif kind == "list":
                self.list_items(b[1])
            elif kind == "table":
                self.table(b[1])
            elif kind == "quote":
                p = self.para(b[1], indent=False)
                p.paragraph_format.left_indent = Cm(1.27)
            elif kind == "code":
                p = self.doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Cm(0)
                _set_run_font(p.add_run(b[1]), "Consolas", max(8, self.size - 2))
            elif kind == "math":
                p = self.para(b[1], align="center", indent=False, base={"italic": True})
            elif kind == "figure":
                p = self.para(f"[Figure: {b[1]}]", align="center", indent=False, base={"italic": True})
            elif kind == "pagebreak":
                self.page_break()

    # -------------------------------------------------------- front matter
    def front(self, meta):
        kind = self.t["title_page"]
        title = meta.get("title") or "Untitled"
        authors = meta.get("authors") or ""
        affil = meta.get("affiliation") or ""
        abstract = meta.get("abstract") or ""
        kw = meta.get("keywords") or ""
        if isinstance(kw, list):
            kw = ", ".join(kw)
        L = self.L
        if kind == "apa":
            self.blank(3)
            self.para(title, align="center", indent=False, base={"bold": True})
            self.blank(1)
            for line in [authors, affil, meta.get("course"), meta.get("instructor"), meta.get("date")]:
                if line:
                    self.para(line, align="center", indent=False)
            if meta.get("author_note"):
                self.blank(2)
                self.para("Author Note", align="center", indent=False, base={"bold": True})
                self.para(meta["author_note"])
            if abstract:
                self.page_break()
                self.para(L["abstract"], align="center", indent=False, base={"bold": True})
                self.para(abstract, indent=False)
                if kw:
                    p = self.para("", indent=True)
                    _run(p, f"{L['keywords']}: ", {"italic": True})
                    add_inline(p, kw)
            self.page_break()
            self.para(title, align="center", indent=False, base={"bold": True})
        elif kind == "ieee":
            self.para(title, align="center", indent=False, size=24, after=6)
            if authors:
                self.para(authors, align="center", indent=False, size=11)
            if affil:
                self.para(affil, align="center", indent=False, size=10, base={"italic": True}, after=12)
            new = self.doc.add_section(WD_SECTION.CONTINUOUS)
            _set_columns(new, 2)
            if abstract:
                p = self.para("", indent=False, size=9)
                _run(p, f"{L['abstract']}—", {"bold": True, "italic": True, "size": 9})
                add_inline(p, abstract, {"bold": True, "size": 9})
            if kw:
                p = self.para("", indent=False, size=9)
                _run(p, f"{L['index_terms']}—", {"bold": True, "italic": True, "size": 9})
                add_inline(p, kw, {"bold": True, "size": 9})
        elif kind in ("journal", "report"):
            big = 18 if kind == "report" else 16
            self.para(title, align="center" if kind == "journal" else "left", indent=False, size=big,
                      base={"bold": True, **({"color": "1F3864"} if kind == "report" else {})}, after=8)
            if authors:
                self.para(authors, align="center" if kind == "journal" else "left", indent=False, after=2)
            if affil:
                self.para(affil, align="center" if kind == "journal" else "left", indent=False,
                          size=max(8, self.size - 1), base={"italic": True}, after=2)
            if meta.get("date"):
                self.para(meta["date"], align="center" if kind == "journal" else "left", indent=False,
                          size=max(8, self.size - 1))
            if abstract:
                self.blank(1)
                self.para(L["abstract"], indent=False, base={"bold": True}, after=4)
                p = self.para(abstract, indent=False, size=max(8, self.size - 1))
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.right_indent = Cm(0.8)
            if kw:
                p = self.para("", indent=False, size=max(8, self.size - 1))
                p.paragraph_format.left_indent = Cm(0.8)
                _run(p, f"{L['keywords']}: ", {"bold": True, "size": max(8, self.size - 1)})
                add_inline(p, kw, {"size": max(8, self.size - 1)})
            self.blank(1)
            if kind == "report" and self.o.get("toc", True):
                self._toc()
        elif kind == "thesis":
            self.blank(2)
            self.para(title.upper(), align="center", indent=False, size=14, base={"bold": True}, spacing=1.15)
            self.blank(2)
            if meta.get("degree"):
                self.para(meta["degree"], align="center", indent=False, base={"bold": True})
                self.blank(2)
            if authors:
                self.para("Disusun oleh:" if self.lang == "id" else "By:", align="center", indent=False)
                self.para(authors, align="center", indent=False, base={"bold": True})
                if meta.get("student_id"):
                    self.para(meta["student_id"], align="center", indent=False)
            self.blank(5)
            for line in [meta.get("department"), affil, meta.get("city"), meta.get("date")]:
                if line:
                    self.para(line.upper(), align="center", indent=False, base={"bold": True})
            if abstract:
                self.page_break()
                self.para(L["abstract"].upper(), align="center", indent=False, base={"bold": True}, after=12)
                self.para(abstract, spacing=1.0)
                if kw:
                    p = self.para("", indent=False)
                    _run(p, f"{L['keywords']}: ", {"bold": True})
                    add_inline(p, kw, {"italic": True})
            if self.o.get("toc", True):
                self.page_break()
                self._toc()

    def _toc(self):
        self.para(self.L["toc"].upper() if self.t["title_page"] == "thesis" else self.L["toc"],
                  align="center" if self.t["title_page"] == "thesis" else "left", indent=False,
                  base={"bold": True}, after=12)
        p = self.doc.add_paragraph()
        _add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose 'Update Field' to build the table of contents.")

    # -------------------------------------------------------- references
    def references(self, entries, style):
        if not entries:
            return
        spec = self.t["ref_heading"]
        title = self.o.get("ref_title") or self.L["references"]
        p = self.doc.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = bool(spec.get("page_break") or self.t["title_page"] == "apa")
        p.paragraph_format.alignment = ALIGN[spec.get("align", "left")]
        txt = title.upper() if spec.get("upper") else title
        r = p.add_run(txt)
        _set_run_font(r, size=spec.get("size"), bold=bool(spec.get("bold")), small_caps=bool(spec.get("small_caps")),
                      color=spec.get("color") or "000000")
        numeric = STYLES.get(style, {}).get("numeric")
        size = self.t["ref_size"]
        for e in entries:
            q = self.doc.add_paragraph()
            pf = q.paragraph_format
            pf.alignment = ALIGN["left"]
            if self.t["ref_spacing"]:
                pf.line_spacing = self.t["ref_spacing"]
                pf.space_after = Pt(4 if self.t["ref_spacing"] < 2 else 0)
            if numeric:
                m = re.match(r"^(\[\d+\]|\d+\.)\s+(.*)$", e)
                label, rest = (m.group(1), m.group(2)) if m else ("", e)
                pf.left_indent = Cm(0.9)
                pf.first_line_indent = Cm(-0.9)
                pf.tab_stops.add_tab_stop(Cm(0.9))
                _run(q, label + "\t", {"size": size} if size else {})
                add_inline(q, rest, {"size": size} if size else {})
            else:
                pf.left_indent = Cm(1.27)
                pf.first_line_indent = Cm(-1.27)
                add_inline(q, e, {"size": size} if size else {})

    def save(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def build_docx(markdown: str, meta: dict, library: dict, template="apa7", style=None, options=None):
    """Returns (docx_bytes, info dict)."""
    tpl = TEMPLATES.get(template, TEMPLATES["apa7"])
    opts = dict(options or {})
    style = style or tpl["style"]
    meta = dict(meta or {})

    text, cited, missing = process_citations(markdown or "", library, style)
    blocks = parse_blocks(text)
    blocks, _ = _strip_sections(blocks, REF_HEADINGS)
    blocks, abstract_text = _strip_sections(blocks, ABSTRACT_HEADINGS)
    if abstract_text and not meta.get("abstract"):
        meta["abstract"] = abstract_text
    if meta.get("abstract"):
        meta["abstract"] = process_citations(meta["abstract"], library, style)[0]
    # first lone H1 = document title; promote remaining headings one level
    h1 = [b for b in blocks if b[0] == "heading" and b[1] == 1]
    if blocks and blocks[0][0] == "heading" and blocks[0][1] == 1 and len(h1) == 1:
        if not meta.get("title"):
            meta["title"] = blocks[0][2]
        blocks = [("heading", max(1, b[1] - 1), b[2]) if b[0] == "heading" else b for b in blocks[1:]]
    elif blocks and not h1 and any(b[0] == "heading" for b in blocks):
        top = min(b[1] for b in blocks if b[0] == "heading")
        blocks = [("heading", b[1] - top + 1, b[2]) if b[0] == "heading" else b for b in blocks]
    if meta.get("keywords") and isinstance(meta["keywords"], str):
        meta["keywords"] = meta["keywords"].strip()

    b = Builder(tpl, opts)
    b.front(meta)
    b.body(blocks)
    entries = bibliography(cited, style)
    if opts.get("include_uncited"):
        cited_keys = {p["key"] for p in cited}
        extra = [p for k, p in library.items() if k not in cited_keys]
        entries = bibliography(cited + extra, style)
    b.references(entries, style)
    return b.save(), {"cited": len(cited), "missing": missing}
