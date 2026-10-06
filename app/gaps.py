"""Manuscript analysis ("what's next for my article") and Gap Finder.

Manuscript: an uploaded in-progress PDF/DOCX/MD is converted to Markdown; DOIs in it are extracted for
one-click import; uncited sentences are pre-detected without AI; 1 fast-model call returns a structured
diagnosis (topic, missing parts, claims that need sources, search topics, next steps).

Gap Finder: 1 fast call breaks the title/proposal into facets with synonyms; a coverage matrix of
published-paper counts per facet pair comes from OpenAlex (no AI); the closest prior work is collected
with the Discover engine and summarised into an evidence table (1 fast call); the gap / novelty /
research-question text is then written with citations by the writer model (streamed).
"""
import asyncio
import io
import itertools
import json
import re

from . import db
from .discovery import _bool_query, discover
from .llm import complete_json
from .search import _client, make_key, pdf_bytes_to_text

SYS = "You are a senior academic editor and research methodologist."


# ---------------------------------------------------------------- manuscript text
def docx_to_markdown(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    out = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        m = re.match(r"heading (\d)", style)
        if m:
            out.append("#" * min(int(m.group(1)), 4) + " " + t)
        elif style == "title":
            out.append("# " + t)
        elif "list" in style:
            out.append("- " + t)
        else:
            out.append(t)
    for table in doc.tables:
        rows = [[c.text.strip().replace("|", "/") for c in r.cells] for r in table.rows]
        if rows:
            out.append("| " + " | ".join(rows[0]) + " |\n|" + "---|" * len(rows[0]) + "\n" +
                       "\n".join("| " + " | ".join(r) + " |" for r in rows[1:]))
    return "\n\n".join(out)


def manuscript_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return pdf_bytes_to_text(data, max_pages=80)
    if name.endswith(".docx"):
        return docx_to_markdown(data)
    return data.decode("utf-8", "ignore")


DOI_RE = re.compile(r"\b(10\.\d{4,9}/[^\s\"<>,;]+[^\s\"<>,;.)\]])", re.I)
CITED_RE = re.compile(r"\[[\d,\s–-]+\]|\([A-Z][A-Za-z'\-]+(?: et al\.?| and [A-Z][a-z]+| & [A-Z][a-z]+)?,? \d{4}[a-z]?(?:[;,][^)]*)?\)|\[@[^\]]+\]")
CLAIM_HINT = re.compile(r"\b(shown|showed|demonstrat|reported|found|known|widely|increas|decreas|improv|reduc|enhanc|"
                        r"significant|effective|previous|studies|has been|have been|is used|are used|\d+(\.\d+)?\s?%)", re.I)


def uncited_claims(text: str, limit=40) -> list[str]:
    body = re.split(r"\n#+\s*(references|bibliography|daftar pustaka)\b", text, flags=re.I)[0]
    out = []
    for para in re.split(r"\n\s*\n", body):
        if para.lstrip().startswith(("#", "|")):
            continue
        for sent in re.split(r"(?<=[.!?])\s+(?=[A-Z])", para.replace("\n", " ")):
            s = sent.strip()
            if 60 < len(s) < 400 and not CITED_RE.search(s) and CLAIM_HINT.search(s):
                out.append(s)
    return out[:limit]


async def analyze_manuscript(text: str, paper_type: str = "") -> dict:
    claims = uncited_claims(text)
    dois = list(dict.fromkeys(d.rstrip(".") for d in DOI_RE.findall(text)))
    headings = re.findall(r"^#+\s+(.+)$", text, re.M)
    words = len(re.findall(r"\b\w+\b", text))
    excerpt = text[:24000]  # ~6k tokens: title, abstract, intro and the start of methods carry the plan
    data = await complete_json(
        f"Manuscript in progress ({words} words; headings: {headings[:40]}).\n<manuscript>\n{excerpt}\n</manuscript>\n\n"
        f"Sentences that look like claims without a citation:\n" + "\n".join(f"{i + 1}. {c}" for i, c in enumerate(claims)) + "\n\n"
        "Diagnose this manuscript so the author can finish it. Return JSON:\n"
        '{"title":"working title","research_question":"","field":"","paper_type":"","stage":"early draft|partial draft|near complete",'
        '"summary":"3 sentences","missing_sections":[""],"weak_sections":[{"section":"","issue":""}],'
        '"claims_needing_sources":[{"n":1,"why":""}],'
        '"search_topics":[{"topic":"short search phrase","why":"what it would support"}],'
        '"next_steps":["concrete, ordered actions to finish the article"]}\n'
        "claims_needing_sources: pick from the numbered list only the claims that truly need a source.",
        system=SYS, fast=True, task="manuscript-analyze", max_tokens=2500)
    picked = []
    for c in data.get("claims_needing_sources", []):
        try:
            picked.append({"sentence": claims[int(c["n"]) - 1], "why": c.get("why", "")})
        except (KeyError, ValueError, IndexError, TypeError):
            pass
    data["claims_needing_sources"] = picked
    data.update(words=words, dois=dois[:150], headings=headings[:60])
    return data


# ---------------------------------------------------------------- Gap Finder
async def facets(topic: str) -> dict:
    data = await complete_json(
        f"Research title / proposal:\n<topic>{topic}</topic>\n\n"
        "Break it into 3-6 FACETS of the research context, each with a role (system/material, variable/intervention, "
        "outcome/property, method, population/setting, application) and the synonyms, abbreviations and technical "
        "names used in papers. Also state the proposed contribution in one sentence.\n"
        'Return JSON: {"facets":[{"name":"","role":"","synonyms":[""],"essential":true}],"contribution":""}',
        system=SYS, fast=True, task="gaps-facets", max_tokens=1200)
    fs = [f for f in data.get("facets", []) if f.get("name")][:6]
    for f in fs:
        f["synonyms"] = list(dict.fromkeys([f["name"]] + [x for x in f.get("synonyms", []) if x]))[:8]
    return {"facets": fs, "contribution": data.get("contribution", "")}


async def _count(q: str, year_from=None) -> int:
    params = {"search": q, "per-page": 1, "select": "id"}
    if year_from:
        params["filter"] = f"from_publication_date:{year_from}-01-01"
    async with _client() as c:
        r = await c.get("https://api.openalex.org/works", params=params)
        r.raise_for_status()
        return r.json()["meta"]["count"]


async def coverage(fs: list[dict], recent_from: int | None = None) -> dict:
    """Published-paper counts for every facet, every facet pair and all facets together (OpenAlex)."""
    jobs = {f["name"]: _bool_query([f]) for f in fs}
    for a, b in itertools.combinations(fs, 2):
        jobs[f"{a['name']} × {b['name']}"] = _bool_query([a, b])
    if len(fs) > 2:
        jobs["ALL facets"] = _bool_query(fs)
    keys = list(jobs)
    counts = await asyncio.gather(*(_count(jobs[k]) for k in keys), return_exceptions=True)
    recent = await asyncio.gather(*(_count(jobs[k], recent_from) for k in keys), return_exceptions=True) if recent_from else [None] * len(keys)
    out = {}
    for k, c, r in zip(keys, counts, recent):
        out[k] = {"total": c if isinstance(c, int) else None, "recent": r if isinstance(r, int) else None}
    return out


async def prior_work(topic: str, sources: list[str], n: int = 25) -> dict:
    found = await discover(topic, sources, snowball_seeds=3, max_results=60)
    top = [p for p in found["results"] if p.get("abstract")][:n]
    if not top:
        return {"papers": [], "table": [], "plan": found["plan"]}
    listing = "\n\n".join(f"[{i + 1}] {p['title']} ({p.get('year')})\n{p['abstract'][:650]}" for i, p in enumerate(top))
    data = await complete_json(
        f"Research topic: {topic}\n\nPapers:\n{listing}\n\nFor each paper extract, from its abstract only, a compact row: "
        "system/material, variables studied, method, main outcome/finding (with numbers if given), and limitation or what it "
        'did NOT study (if inferable). Return JSON: {"rows":[{"id":1,"system":"","variables":"","method":"","finding":"","not_studied":""}]}',
        system=SYS, fast=True, task="gaps-extract", max_tokens=4000)
    rows = []
    for r in data.get("rows", []):
        try:
            p = top[int(r["id"]) - 1]
        except (KeyError, ValueError, IndexError, TypeError):
            continue
        rows.append({**{k: r.get(k, "") for k in ("system", "variables", "method", "finding", "not_studied")},
                     "title": p["title"], "year": p.get("year"), "idx": int(r["id"]) - 1})
    return {"papers": top, "table": rows, "plan": found["plan"], "total_found": found["total_found"]}


def assign_keys(project_id: str, papers: list[dict]):
    lib = {p["key"] for p in db.list_papers(project_id)}
    used = set()
    for p in papers:
        existing = db.find_existing(project_id, p)
        if existing:
            p["key"] = existing["key"]
        else:
            base = k = make_key(p)
            i = 0
            while k in lib or k in used:
                i += 1
                k = base + "abcdefghijklmnopqrstuvwxyz"[(i - 1) % 26]
            p["key"] = k
        used.add(p["key"])


def p_gap_writeup(topic, fs, contribution, cov, papers, table, language="English"):
    cov_txt = "\n".join(f"- {k}: {v['total']} papers" + (f" ({v['recent']} in recent years)" if v.get("recent") is not None else "")
                        for k, v in cov.items())
    rows = "\n".join(f"- @{papers[r['idx']]['key']} ({r['year']}): system={r['system']}; variables={r['variables']}; "
                     f"method={r['method']}; finding={r['finding']}; not studied={r['not_studied']}" for r in table)
    return (f"Proposed research:\n<topic>{topic}</topic>\nIntended contribution: {contribution}\n\n"
            f"Facets: {json.dumps([{k: f[k] for k in ('name', 'role')} for f in fs], ensure_ascii=False)}\n\n"
            f"Coverage in the published literature (OpenAlex counts for facet combinations):\n{cov_txt}\n\n"
            f"Closest prior work (evidence table; cite ONLY these keys):\n{rows}\n\n"
            f"Write in {language}, Markdown, academic register:\n"
            "## State of the art: what is already established (synthesise, cite heavily)\n"
            "## Research gaps: numbered '### Gap n: <name>' items. Each gap must be supported by 3-5 citations showing what "
            "has been done and must explain precisely what is missing, using the coverage counts as evidence where useful "
            "(e.g. few or no papers combining two facets).\n"
            "## Novelty statement: 1 paragraph stating what is new in the proposed research versus the closest works.\n"
            "## Closest existing works & overlap risk: the 3-5 most similar papers and how the proposal differs.\n"
            "## Research questions / hypotheses: 2-4 items.\n"
            "## Suggested Introduction paragraph: a ready-to-use gap → aim paragraph with citations.\n"
            "Cite with [@key] / [@key1; @key2]. Never invent keys, counts or findings. Be honest if the idea seems "
            "already well covered, and suggest how to sharpen it.")
